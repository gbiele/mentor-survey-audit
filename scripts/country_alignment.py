#!/usr/bin/env python3
"""Align independent-country surveys to the canonical English core via translation and review."""

from __future__ import annotations

import json
import os
import re
import unicodedata
from collections import defaultdict
from copy import deepcopy
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
MAPS_DIR = ROOT / "data" / "country_maps"
TRANSLATION_CACHE = MAPS_DIR / "translation_cache.json"

ID_OVERLAP_THRESHOLD = 0.4
BATCH_SIZE = 20
FALLBACK_MATCH_THRESHOLD = 0.58
FALLBACK_MATCH_GAP = 0.07
INFER_SCALE_MATCH_THRESHOLD = 0.42
STEM_ONLY_OPTION_MISMATCH_THRESHOLD = 0.40
EXACT_NAME_STEM_THRESHOLD = 0.30

# Demographics / single-stem items where country option lists often differ (ESO grades, etc.)
OPTION_MISMATCH_TOLERANT_PREFIXES = frozenset(
    {
        "grade",
        "edumom",
        "edudad",
        "ladder",
        "bage",
        "bmborn",
        "consent",
        "byear",
        "bmonth",
        "born",
        "gender",
    }
)

SOURCE_LANG = {
    "spain": "es",
    "germany": "en",
    "france": "fr",
    "italy": "it",
}

REVIEW_COLUMNS = [
    "country_orig_id",
    "section",
    "country_stem",
    "country_item",
    "country_options",
    "translated_text",
    "canonical_variable",
    "canonical_stem",
    "canonical_item",
    "canonical_options",
    "status",
    "notes",
]

STATUS_FILLS = {
    "matched": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "probable": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "conflict": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "extra": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    "missing_core": PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid"),
}

COVERED_STATUSES = frozenset({"matched", "probable", "conflict"})


@dataclass
class TranslationResult:
    orig_variable: str
    translated_stem: str = ""
    translated_item: str = ""
    translated_options: list[str] = field(default_factory=list)
    canonical_id: str | None = None
    canonical_stem: str = ""
    canonical_item: str = ""
    canonical_options: list[str] = field(default_factory=list)


@dataclass
class AlignmentRow:
    country_orig_id: str
    section: str
    country_stem: str
    country_item: str
    country_options: str
    translated_text: str
    canonical_variable: str
    canonical_stem: str
    canonical_item: str
    canonical_options: str
    status: str
    notes: str = ""


def default_source_lang(source_name: str) -> str:
    return SOURCE_LANG.get(source_name.lower(), "en")


def _option_labels(options: list[Any]) -> list[str]:
    return [getattr(o, "label", str(o)) for o in (options or [])]


def _join_options(options: list[str]) -> str:
    return " | ".join(o for o in options if o)


def flatten_country_items(codebook: list[Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for q in codebook:
        for it in q.items:
            rows.append(
                {
                    "orig_variable": it.variable,
                    "section": q.section,
                    "country_stem": q.stem,
                    "country_item": it.item_text or "",
                    "country_options": _option_labels(q.options),
                    "question": q,
                    "item": it,
                }
            )
    return rows


def build_core_whitelist(core_codebook: list[Any]) -> list[dict[str, Any]]:
    whitelist: list[dict[str, Any]] = []
    for q in core_codebook:
        for it in q.items:
            whitelist.append(
                {
                    "variable": it.short_name or it.variable,
                    "section": q.section,
                    "question_stem": q.stem,
                    "item_text": it.item_text or "",
                    "options": _option_labels(q.options),
                }
            )
    return whitelist


def core_orig_ids(core_codebook: list[Any]) -> set[str]:
    return {it.variable for q in core_codebook for it in q.items}


def country_orig_ids(country_codebook: list[Any]) -> set[str]:
    return {it.variable for q in country_codebook for it in q.items}


def id_overlap_ratio(country_codebook: list[Any], core_codebook: list[Any]) -> float:
    country_ids = country_orig_ids(country_codebook)
    if not country_ids:
        return 0.0
    core_ids = core_orig_ids(core_codebook)
    return len(country_ids & core_ids) / len(country_ids)


def _cache_key(source_name: str, orig_variable: str) -> str:
    return f"{source_name}:{orig_variable}"


def _load_translation_cache(cache_file: Path) -> dict[str, dict[str, Any]]:
    if not cache_file.exists():
        return {}
    try:
        with open(cache_file, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def _save_translation_cache(cache_file: Path, cache: dict[str, dict[str, Any]]) -> None:
    cache_file.parent.mkdir(parents=True, exist_ok=True)
    with open(cache_file, "w", encoding="utf-8") as f:
        json.dump(cache, f, indent=2, ensure_ascii=False)


def _parse_translation_entry(entry: dict[str, Any]) -> TranslationResult:
    opts = entry.get("translated_options") or entry.get("canonical_options") or []
    canon_opts = entry.get("canonical_options") or []
    if isinstance(opts, str):
        opts = [x.strip() for x in opts.split("|") if x.strip()]
    if isinstance(canon_opts, str):
        canon_opts = [x.strip() for x in canon_opts.split("|") if x.strip()]
    canonical_id = entry.get("canonical_id")
    if canonical_id in ("", "null", None):
        canonical_id = None
    return TranslationResult(
        orig_variable=str(entry.get("orig_variable", "")),
        translated_stem=str(entry.get("translated_stem", "")),
        translated_item=str(entry.get("translated_item", "")),
        translated_options=list(opts),
        canonical_id=canonical_id,
        canonical_stem=str(entry.get("canonical_stem", "")),
        canonical_item=str(entry.get("canonical_item", "")),
        canonical_options=list(canon_opts),
    )


def translate_country_texts(
    country_codebook: list[Any],
    source_lang: str,
    core_whitelist: list[dict[str, Any]],
    source_name: str,
    cache_file: Path = TRANSLATION_CACHE,
    force_refresh: bool = False,
) -> dict[str, TranslationResult]:
    """Translate country items with whitelist-guided canonical snapping."""
    rows = flatten_country_items(country_codebook)
    if source_lang.lower() in ("en", "english") or not rows:
        return {}

    cache = _load_translation_cache(cache_file)
    results: dict[str, TranslationResult] = {}
    pending: list[dict[str, Any]] = []

    for row in rows:
        key = _cache_key(source_name, row["orig_variable"])
        if not force_refresh and key in cache:
            results[row["orig_variable"]] = _parse_translation_entry(cache[key])
        else:
            pending.append(row)

    api_key = os.environ.get("GEMINI_API_KEY")
    if pending and api_key:
        try:
            from google import genai
            from google.genai import types

            client = genai.Client(api_key=api_key)
            for i in range(0, len(pending), BATCH_SIZE):
                batch = pending[i : i + BATCH_SIZE]
                batch_payload = []
                for row in batch:
                    batch_payload.append(
                        {
                            "orig_variable": row["orig_variable"],
                            "section": row["section"],
                            "question_stem": row["country_stem"],
                            "item_text": row["country_item"],
                            "options": row["country_options"],
                        }
                    )

                prompt = f"""
You are aligning a youth mental health survey from language "{source_lang}" to a canonical English core.

For each country item below, compare it to the canonical English whitelist.
Return ONLY valid JSON as an array of objects with these fields:
- orig_variable
- translated_stem, translated_item, translated_options (your English translation)
- canonical_id (whitelist variable name if clearly the same question, else null)
- canonical_stem, canonical_item, canonical_options (copy whitelist text verbatim when canonical_id is set; else empty strings)

Rules:
- Use whitelist text when the country item is clearly the same question (match item + options, not stem alone).
- If only related or country-specific, set canonical_id to null. Do not pick the closest item.
- Never overwrite meaning; keep translations faithful.

Country items:
{json.dumps(batch_payload, ensure_ascii=False, indent=2)}

Canonical English whitelist:
{json.dumps(core_whitelist, ensure_ascii=False, indent=2)}
"""
                response = client.models.generate_content(
                    model="gemini-3.7-flash",
                    contents=prompt,
                    config=types.GenerateContentConfig(
                        response_mime_type="application/json",
                        temperature=0.1,
                    ),
                )
                parsed = json.loads(response.text)
                entries = parsed if isinstance(parsed, list) else parsed.get("items", [])
                for entry in entries:
                    tr = _parse_translation_entry(entry)
                    if tr.orig_variable:
                        key = _cache_key(source_name, tr.orig_variable)
                        cache[key] = {
                            "orig_variable": tr.orig_variable,
                            "translated_stem": tr.translated_stem,
                            "translated_item": tr.translated_item,
                            "translated_options": tr.translated_options,
                            "canonical_id": tr.canonical_id,
                            "canonical_stem": tr.canonical_stem,
                            "canonical_item": tr.canonical_item,
                            "canonical_options": tr.canonical_options,
                        }
                        results[tr.orig_variable] = tr
                print(
                    f"  Translated batch {i // BATCH_SIZE + 1}/"
                    f"{(len(pending) + BATCH_SIZE - 1) // BATCH_SIZE} ({len(batch)} items).",
                    flush=True,
                )
            _save_translation_cache(cache_file, cache)
        except Exception as exc:
            print(f"Warning: Gemini translation failed ({exc}); using cache only.", flush=True)
    elif pending and not api_key:
        print(
            f"Note: {len(pending)} country items need translation but GEMINI_API_KEY is not set.",
            flush=True,
        )

    return results


def _core_lookup(core_codebook: list[Any]) -> dict[str, tuple[Any, Any]]:
    lookup: dict[str, tuple[Any, Any]] = {}
    for q in core_codebook:
        for it in q.items:
            lookup[it.short_name or it.variable] = (q, it)
    return lookup


def _core_orig_to_canonical(core_codebook: list[Any]) -> dict[str, str]:
    mapping: dict[str, str] = {}
    for q in core_codebook:
        for it in q.items:
            mapping[it.variable] = it.short_name or it.variable
    return mapping


def _normalize_text(text: str) -> str:
    if not text:
        return ""
    s = unicodedata.normalize("NFKC", str(text).casefold())
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _text_similarity(a: str, b: str) -> float:
    a_n = _normalize_text(a)
    b_n = _normalize_text(b)
    if not a_n or not b_n:
        return 0.0
    if a_n == b_n:
        return 1.0
    if a_n in b_n or b_n in a_n:
        return 0.9
    return SequenceMatcher(None, a_n, b_n).ratio()


def _options_compatible(country_options: list[str], core_options: list[str]) -> bool:
    if not country_options or not core_options:
        return True
    return abs(len(country_options) - len(core_options)) <= 1


def _variable_prefix(variable: str) -> str:
    prefix = _scale_prefix(variable)
    if prefix:
        return prefix
    m = re.match(r"^(.+?)\d+$", variable)
    return m.group(1).lower() if m else variable.lower()


def _allows_option_count_mismatch(variable: str) -> bool:
    return _variable_prefix(variable) in OPTION_MISMATCH_TOLERANT_PREFIXES


def _split_stem_item(stem: str, item: str) -> tuple[str, str]:
    """When item text is embedded after a colon in the stem, split it out."""
    if item:
        return stem, item
    if ":" in stem:
        left, right = stem.split(":", 1)
        if right.strip():
            return left.strip(), right.strip()
    return stem, item


def _effective_country_item(
    translation: TranslationResult | None,
    country_item: str,
) -> str:
    if translation and translation.translated_item:
        return translation.translated_item
    return country_item


def _stem_and_item_scores(
    translation: TranslationResult | None,
    country_stem: str,
    country_item: str,
    core_stem: str,
    core_item: str,
) -> tuple[float, float]:
    country_stem, country_item = _split_stem_item(country_stem, country_item)
    translated_stem = ""
    translated_item = ""
    if translation:
        translated_stem, translated_item = _split_stem_item(
            translation.translated_stem,
            translation.translated_item,
        )

    stem_scores: list[float] = []
    item_scores: list[float] = []

    if translated_stem and core_stem:
        stem_scores.append(_text_similarity(translated_stem, core_stem))
    if country_stem and core_stem:
        stem_scores.append(_text_similarity(country_stem, core_stem) * 0.85)

    effective_item = translated_item or country_item
    if effective_item and core_item:
        item_scores.append(_text_similarity(effective_item, core_item))
    if (
        country_item
        and core_item
        and country_item != effective_item
    ):
        item_scores.append(_text_similarity(country_item, core_item) * 0.75)

    stem_score = max(stem_scores) if stem_scores else 0.0
    item_score = max(item_scores) if item_scores else 0.0
    return stem_score, item_score


def _parent_education_hint(stem: str, item: str, translation: TranslationResult | None = None) -> str | None:
    parts = [stem, item]
    if translation:
        parts.extend([translation.translated_stem, translation.translated_item])
    text = _normalize_text(" ".join(p for p in parts if p))
    if re.search(r"\b(mother|madre|mom|mama)\b", text):
        return "edumom1"
    if re.search(r"\b(father|padre|dad|papa)\b", text):
        return "edudad1"
    return None


def _ladder_question_hint(
    stem: str,
    item: str,
    translation: TranslationResult | None = None,
) -> bool:
    parts = [stem, item]
    if translation:
        parts.extend([translation.translated_stem, translation.translated_item])
    text = _normalize_text(" ".join(p for p in parts if p))
    return bool(re.search(r"\b(ladder|escalera)\b", text))


def _immigration_question_hint(stem: str, item: str) -> str | None:
    """Disambiguate combined age-at-arrival vs birth-country questions."""
    text = _normalize_text(f"{stem} {item}")
    age_arrival = bool(
        re.search(
            r"(age|edad|anos|years).*(arriv|lleg|immig|migr|move|pais|country|when you came)",
            text,
        )
    )
    birth_country = bool(
        re.search(r"(country|pais|naci).*(born|birth|nac|origin)", text)
    )
    if age_arrival and not birth_country:
        return "bage1"
    if birth_country and not age_arrival:
        return "bmborn1"
    return None


def _match_score(
    translation: TranslationResult | None,
    country_stem: str,
    country_item: str,
    country_options: list[str],
    core_entry: dict[str, Any],
) -> float:
    core_stem = core_entry.get("question_stem", "")
    core_item = core_entry.get("item_text", "")
    core_options = core_entry.get("options", [])
    variable = core_entry.get("variable", "")

    stem_score, item_score = _stem_and_item_scores(
        translation, country_stem, country_item, core_stem, core_item
    )
    options_ok = _options_compatible(country_options, core_options)

    if not options_ok:
        if _allows_option_count_mismatch(variable):
            if stem_score >= STEM_ONLY_OPTION_MISMATCH_THRESHOLD:
                return min(1.0, max(stem_score, item_score * 0.95))
            return 0.0
        if core_item:
            # Matrix row: require item-level agreement when option counts diverge.
            if item_score >= 0.42:
                return min(1.0, max(item_score, stem_score * 0.65))
            return 0.0
        if stem_score >= 0.62:
            return min(1.0, stem_score * 0.92)
        return 0.0

    scores: list[float] = []
    if item_score:
        scores.append(item_score * (1.2 if core_item else 1.0))
    if stem_score:
        scores.append(stem_score)

    if translation and translation.translated_options and core_options:
        opt_scores = [
            _text_similarity(t_opt, c_opt)
            for t_opt, c_opt in zip(translation.translated_options[:6], core_options[:6])
        ]
        if opt_scores:
            scores.append(sum(opt_scores) / len(opt_scores) * 0.8)

    if not scores:
        return 0.0
    return min(1.0, max(scores))


def _translation_from_row(row: AlignmentRow) -> TranslationResult | None:
    if not row.translated_text:
        return None
    text = row.translated_text
    opts: list[str] = []
    if " | " in text:
        main, opts_part = text.rsplit(" | ", 1)
        opts = [x.strip() for x in opts_part.split("|") if x.strip()]
    else:
        main = text
    if " : " in main:
        stem, item = main.split(" : ", 1)
    else:
        stem, item = main, ""
    return TranslationResult(
        orig_variable=row.country_orig_id,
        translated_stem=stem.strip(),
        translated_item=item.strip(),
        translated_options=opts,
    )


def _row_match_score(row: AlignmentRow, core_entry: dict[str, Any]) -> float:
    country_options = [x.strip() for x in row.country_options.split("|") if x.strip()]
    return _match_score(
        _translation_from_row(row),
        row.country_stem,
        row.country_item,
        country_options,
        core_entry,
    )


def _best_scale_prefix_for_row(
    row: AlignmentRow,
    core_whitelist: list[dict[str, Any]],
) -> tuple[str | None, float]:
    """Pick the scale family that best matches a country row (avoids cross-scale steals)."""
    best_prefix = None
    best_score = 0.0
    seen_prefixes: set[str] = set()
    for core_entry in core_whitelist:
        var = core_entry["variable"]
        prefix = _scale_prefix(var)
        if not prefix or prefix in seen_prefixes:
            continue
        seen_prefixes.add(prefix)
        target = _first_scale_variable(prefix, core_whitelist)
        if not target:
            continue
        target_entry = next(c for c in core_whitelist if c["variable"] == target)
        score = _row_match_score(row, target_entry)
        if score > best_score:
            best_score = score
            best_prefix = prefix
    return best_prefix, best_score


def _shared_id_blocks_expected_match(
    orig: str,
    canonical_variable: str,
    translation: TranslationResult | None,
    core_lookup: dict[str, tuple[Any, Any]],
    core_orig_map: dict[str, str],
    core_whitelist: list[dict[str, Any]],
) -> tuple[bool, str]:
    """Reject only when a shared EUSurvey ID is mapped to the colliding core owner with weak text."""
    if orig not in core_orig_map or not canonical_variable:
        return False, ""
    expected = core_orig_map[orig]
    if canonical_variable != expected:
        return False, ""
    core_entry = next((c for c in core_whitelist if c["variable"] == canonical_variable), None)
    if not core_entry:
        return True, f"Shared EUSurvey ID {orig} maps to core {expected}; rejected"
    score = _match_score(
        translation,
        translation.translated_stem if translation else "",
        translation.translated_item if translation else "",
        translation.translated_options if translation else [],
        core_entry,
    )
    if score < 0.45:
        return True, (
            f"Shared EUSurvey ID {orig} rejected for core {expected} "
            f"(similarity {score:.2f})"
        )
    return False, ""


def _exact_variable_name_match(
    orig: str,
    translation: TranslationResult | None,
    row: dict[str, Any],
    core_whitelist: list[dict[str, Any]],
    assigned: set[str],
) -> tuple[str, str]:
    """Match when country orig_variable equals a core short name (e.g. ks2 -> ks2)."""
    for core_entry in core_whitelist:
        var = core_entry["variable"]
        if var.lower() != orig.lower() or var in assigned:
            continue
        stem_score, item_score = _stem_and_item_scores(
            translation,
            row["country_stem"],
            row["country_item"],
            core_entry.get("question_stem", ""),
            core_entry.get("item_text", ""),
        )
        if max(stem_score, item_score) >= EXACT_NAME_STEM_THRESHOLD:
            return var, f"Exact variable name match ({max(stem_score, item_score):.2f})"
    return "", ""


def fallback_match_canonical(
    translation: TranslationResult | None,
    row: dict[str, Any],
    core_whitelist: list[dict[str, Any]],
    assigned: set[str],
    core_orig_map: dict[str, str],
    orig: str,
) -> tuple[str, str]:
    """Deterministic match when LLM leaves canonical_id empty."""
    exact_var, exact_note = _exact_variable_name_match(
        orig, translation, row, core_whitelist, assigned
    )
    if exact_var:
        return exact_var, exact_note

    best_var = ""
    best_score = 0.0
    second_score = 0.0
    hint = _immigration_question_hint(row["country_stem"], row["country_item"])
    parent_hint = _parent_education_hint(
        row["country_stem"], row["country_item"], translation
    )
    ladder_hint = _ladder_question_hint(
        row["country_stem"], row["country_item"], translation
    )

    for core_entry in core_whitelist:
        var = core_entry["variable"]
        if var in assigned:
            continue
        expected = core_orig_map.get(orig)
        if expected and var == expected:
            continue
        score = _match_score(
            translation,
            row["country_stem"],
            row["country_item"],
            row["country_options"],
            core_entry,
        )
        if hint and var == hint:
            score = min(1.0, score + 0.12)
        elif hint and var in {"bage1", "bmborn1"} and var != hint:
            score *= 0.55
        if parent_hint and var == parent_hint:
            score = min(1.0, max(score, 0.72))
        elif parent_hint and var in {"edumom1", "edudad1"} and var != parent_hint:
            score *= 0.35
        if ladder_hint and var == "ladder1":
            score = min(1.0, max(score, 0.72))
        if score > best_score:
            second_score = best_score
            best_score = score
            best_var = var
        elif score > second_score:
            second_score = score

    threshold = FALLBACK_MATCH_THRESHOLD
    gap = FALLBACK_MATCH_GAP
    if best_var and _allows_option_count_mismatch(best_var):
        threshold = STEM_ONLY_OPTION_MISMATCH_THRESHOLD
        gap = 0.04

    if (
        best_var
        and best_score >= threshold
        and (best_score - second_score) >= gap
    ):
        return best_var, f"Fallback text match ({best_score:.2f})"
    return "", ""


def _scale_prefix(variable: str) -> str | None:
    m = re.match(
        r"^(bcfpi_[a-z]+|ks|cyrm|ace|rpginsfy|tnqactmh|tnqaomhs|dlup|ewcvtqaw|fetnqact|smharm|smwd|smwda)\d+$",
        variable,
        re.I,
    )
    if m:
        return re.sub(r"\d+$", "", variable.lower())
    m = re.match(r"^(.+?)\d+$", variable)
    return m.group(1).lower() if m else None


def _stem_group_key(stem: str) -> str:
    base = stem.split(":")[0] if ":" in stem else stem
    return _normalize_text(base)[:80]


def _first_scale_variable(prefix: str, core_whitelist: list[dict[str, Any]]) -> str | None:
    candidates = [
        c["variable"]
        for c in core_whitelist
        if re.match(rf"^{re.escape(prefix)}\d+$", c["variable"], re.I)
    ]
    if not candidates:
        return None
    return min(candidates, key=lambda v: int(re.search(r"(\d+)$", v).group(1)))


def _infer_scale_row_score(row: AlignmentRow, target_entry: dict[str, Any]) -> float:
    score = _row_match_score(row, target_entry)
    if score >= INFER_SCALE_MATCH_THRESHOLD:
        return score
    stem, item = _split_stem_item(row.country_stem, row.country_item)
    tr = _translation_from_row(row)
    if item or (tr and tr.translated_item):
        return score
    stem_score, _ = _stem_and_item_scores(
        tr,
        stem,
        item,
        target_entry.get("question_stem", ""),
        target_entry.get("item_text", ""),
    )
    if stem_score >= INFER_SCALE_MATCH_THRESHOLD:
        return stem_score
    # Matrix header row in a stem group that already has sibling scale matches.
    return 0.75


def infer_scale_first_items(
    alignment: list[AlignmentRow],
    core_whitelist: list[dict[str, Any]],
) -> None:
    """If scale items 2..n matched in a stem group, infer the first item."""
    by_stem: dict[str, list[AlignmentRow]] = defaultdict(list)
    for row in alignment:
        by_stem[_stem_group_key(row.country_stem)].append(row)

    assigned = {r.canonical_variable for r in alignment if r.canonical_variable}

    for rows in by_stem.values():
        prefix_nums: dict[str, set[int]] = defaultdict(set)
        prefix_stem_groups: dict[str, set[str]] = defaultdict(set)
        unmatched: list[AlignmentRow] = []
        for row in rows:
            if row.canonical_variable:
                prefix = _scale_prefix(row.canonical_variable)
                num_m = re.search(r"(\d+)$", row.canonical_variable)
                if prefix and num_m:
                    prefix_nums[prefix].add(int(num_m.group(1)))
                    prefix_stem_groups[prefix].add(_stem_group_key(row.country_stem))
            elif row.status == "extra":
                unmatched.append(row)

        for prefix, nums in prefix_nums.items():
            if len(nums) < 2 or 1 in nums or not unmatched:
                continue
            target = _first_scale_variable(prefix, core_whitelist)
            if not target or target in assigned:
                continue
            target_entry = next(c for c in core_whitelist if c["variable"] == target)
            stem_groups = prefix_stem_groups.get(prefix, set())
            best_row = None
            best_score = 0.0
            for candidate in unmatched:
                if _stem_group_key(candidate.country_stem) not in stem_groups:
                    continue
                score = _infer_scale_row_score(candidate, target_entry)
                if score > best_score:
                    best_score = score
                    best_row = candidate
            if not best_row or best_score < INFER_SCALE_MATCH_THRESHOLD:
                continue
            best_row.canonical_variable = target
            best_row.status = "probable"
            best_row.notes = (best_row.notes + "; " if best_row.notes else "") + (
                f"Inferred {target} from sibling scale matches ({prefix}*, {best_score:.2f})"
            )
            assigned.add(target)
            unmatched.remove(best_row)


def _resolve_canonical_for_row(
    row: dict[str, Any],
    translation: TranslationResult | None,
    core_whitelist: list[dict[str, Any]],
    core_lookup: dict[str, tuple[Any, Any]],
    core_orig_map: dict[str, str],
    assigned: set[str],
) -> tuple[str, str]:
    orig = row["orig_variable"]
    canonical_variable = ""
    notes = ""

    if translation and translation.canonical_id:
        canonical_variable = translation.canonical_id
        blocked, block_note = _shared_id_blocks_expected_match(
            orig,
            canonical_variable,
            translation,
            core_lookup,
            core_orig_map,
            core_whitelist,
        )
        if blocked:
            canonical_variable = ""
            notes = block_note

    if not canonical_variable:
        canonical_variable, notes = fallback_match_canonical(
            translation,
            row,
            core_whitelist,
            assigned,
            core_orig_map,
            orig,
        )

    return canonical_variable, notes


def _determine_status(
    canonical_id: str | None,
    canonical_counts: dict[str, int],
    country_options: list[str],
    core_options: list[str],
) -> str:
    if not canonical_id:
        return "extra"
    if canonical_counts.get(canonical_id, 0) > 1:
        return "conflict"
    if country_options and core_options:
        if abs(len(country_options) - len(core_options)) > 1:
            return "probable"
    return "matched"


def _row_from_parts(
    row: dict[str, Any],
    translation: TranslationResult | None,
    canonical_variable: str,
    core_lookup: dict[str, tuple[Any, Any]],
    canonical_counts: dict[str, int],
    status_override: str | None = None,
    notes: str = "",
) -> AlignmentRow:
    tr = translation
    translated_text = ""
    if tr:
        parts = [tr.translated_stem, tr.translated_item]
        translated_text = " : ".join(p for p in parts if p)
        if tr.translated_options:
            translated_text += " | " + _join_options(tr.translated_options)

    canonical_stem = ""
    canonical_item = ""
    canonical_options = ""
    if canonical_variable and canonical_variable in core_lookup:
        cq, cit = core_lookup[canonical_variable]
        canonical_stem = cq.stem
        canonical_item = cit.item_text or ""
        canonical_options = _join_options(_option_labels(cq.options))
    elif tr and tr.canonical_id:
        canonical_stem = tr.canonical_stem
        canonical_item = tr.canonical_item
        canonical_options = _join_options(tr.canonical_options)

    status = status_override or _determine_status(
        canonical_variable or None,
        canonical_counts,
        row["country_options"],
        canonical_options.split(" | ") if canonical_options else [],
    )
    return AlignmentRow(
        country_orig_id=row["orig_variable"],
        section=row["section"],
        country_stem=row["country_stem"],
        country_item=row["country_item"],
        country_options=_join_options(row["country_options"]),
        translated_text=translated_text,
        canonical_variable=canonical_variable,
        canonical_stem=canonical_stem,
        canonical_item=canonical_item,
        canonical_options=canonical_options,
        status=status,
        notes=notes,
    )


def _covered_core_variables(alignment: list[AlignmentRow]) -> set[str]:
    return {
        r.canonical_variable
        for r in alignment
        if r.canonical_variable and r.status in COVERED_STATUSES
    }


def _missing_core_rows(
    core_codebook: list[Any],
    country_alignment: list[AlignmentRow],
) -> list[AlignmentRow]:
    """Core items with no matching country row (shown in red in the review workbook)."""
    covered = _covered_core_variables(country_alignment)
    missing: list[AlignmentRow] = []
    for q in core_codebook:
        for it in q.items:
            var = it.short_name or it.variable
            if var in covered:
                continue
            missing.append(
                AlignmentRow(
                    country_orig_id="",
                    section=q.section,
                    country_stem="",
                    country_item="",
                    country_options="",
                    translated_text="",
                    canonical_variable=var,
                    canonical_stem=q.stem,
                    canonical_item=it.item_text or "",
                    canonical_options=_join_options(_option_labels(q.options)),
                    status="missing_core",
                    notes="Core question not matched to any country item",
                )
            )
    return missing


def _core_variable_order(core_codebook: list[Any]) -> dict[str, int]:
    order: dict[str, int] = {}
    for q in core_codebook:
        for it in q.items:
            order[it.short_name or it.variable] = len(order)
    return order


def _country_variable_order(country_codebook: list[Any]) -> tuple[list[str], dict[str, int]]:
    ordered: list[str] = []
    for q in country_codebook:
        for it in q.items:
            ordered.append(it.variable)
    return ordered, {v: i for i, v in enumerate(ordered)}


def sort_alignment_for_review(
    rows: list[AlignmentRow],
    core_codebook: list[Any],
    country_codebook: list[Any],
) -> list[AlignmentRow]:
    """Order rows by base questionnaire sequence, interleaving country extras by survey position."""
    core_pos = _core_variable_order(core_codebook)
    country_order, country_pos = _country_variable_order(country_codebook)

    orig_to_canon: dict[str, str] = {}
    for row in rows:
        if row.country_orig_id and row.canonical_variable:
            orig_to_canon[row.country_orig_id] = row.canonical_variable

    prev_core: dict[str, int] = {}
    last = -1
    for orig in country_order:
        prev_core[orig] = last
        canon = orig_to_canon.get(orig)
        if canon in core_pos:
            last = core_pos[canon]

    next_core: dict[str, int] = {}
    last = len(core_pos)
    for orig in reversed(country_order):
        next_core[orig] = last
        canon = orig_to_canon.get(orig)
        if canon in core_pos:
            last = core_pos[canon]

    def sort_key(row: AlignmentRow) -> tuple[int, int, int]:
        if row.status == "missing_core":
            pos = core_pos.get(row.canonical_variable, 10_000)
            return (pos, 2, 0)

        orig = row.country_orig_id
        if not orig:
            return (10_001, 0, 0)

        cp = country_pos.get(orig, 10_000)
        canon = row.canonical_variable
        if canon and canon in core_pos:
            return (core_pos[canon], 0, cp)

        pc = prev_core.get(orig, -1)
        nc = next_core.get(orig, len(core_pos))
        if pc < 0:
            return (-1, 1, cp)
        if pc >= nc:
            return (nc, 1, cp)
        return (pc, 1, cp)

    return sorted(rows, key=sort_key)


def _alignment_to_rows(alignment: list[AlignmentRow]) -> list[list[Any]]:
    return [
        [
            r.country_orig_id,
            r.section,
            r.country_stem,
            r.country_item,
            r.country_options,
            r.translated_text,
            r.canonical_variable,
            r.canonical_stem,
            r.canonical_item,
            r.canonical_options,
            r.status,
            r.notes,
        ]
        for r in alignment
    ]


def write_review_workbook(
    review_path: Path,
    country_alignment: list[AlignmentRow],
    core_codebook: list[Any] | None = None,
    country_codebook: list[Any] | None = None,
) -> None:
    review_path.parent.mkdir(parents=True, exist_ok=True)
    rows = list(country_alignment)
    if core_codebook is not None:
        rows.extend(_missing_core_rows(core_codebook, country_alignment))
    if core_codebook is not None and country_codebook is not None:
        rows = sort_alignment_for_review(rows, core_codebook, country_codebook)

    wb = Workbook()
    ws = wb.active
    ws.title = "Review"
    ws.append(REVIEW_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    status_col = REVIEW_COLUMNS.index("status") + 1
    for row_vals in _alignment_to_rows(rows):
        ws.append(row_vals)
        status = row_vals[status_col - 1]
        fill = STATUS_FILLS.get(status)
        if fill:
            for col_idx in range(1, len(REVIEW_COLUMNS) + 1):
                ws.cell(row=ws.max_row, column=col_idx).fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, _ in enumerate(REVIEW_COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 24
    wb.save(review_path)


def load_review_workbook(review_path: Path) -> dict[str, AlignmentRow]:
    wb = load_workbook(review_path, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    rows: dict[str, AlignmentRow] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue

        def val(name: str) -> str:
            i = idx.get(name)
            if i is None:
                return ""
            v = row[i]
            return "" if v is None else str(v).strip()

        orig = val("country_orig_id")
        status = val("status")
        if status == "missing_core" or not orig:
            continue
        rows[orig] = AlignmentRow(
            country_orig_id=orig,
            section=val("section"),
            country_stem=val("country_stem"),
            country_item=val("country_item"),
            country_options=val("country_options"),
            translated_text=val("translated_text"),
            canonical_variable=val("canonical_variable"),
            canonical_stem=val("canonical_stem"),
            canonical_item=val("canonical_item"),
            canonical_options=val("canonical_options"),
            status=val("status") or "extra",
            notes=val("notes"),
        )
    return rows


def align_country_to_core(
    country_codebook: list[Any],
    core_codebook: list[Any],
    review_path: Path,
    source_name: str,
    translations: dict[str, TranslationResult] | None = None,
    force_realign: bool = False,
) -> list[AlignmentRow]:
    """Build or load alignment rows and write/update the Excel review workbook."""
    country_rows = flatten_country_items(country_codebook)
    core_lookup = _core_lookup(core_codebook)
    core_whitelist = build_core_whitelist(core_codebook)
    core_orig_map = _core_orig_to_canonical(core_codebook)
    existing = {} if force_realign or not review_path.exists() else load_review_workbook(review_path)

    alignment: list[AlignmentRow] = []
    seen: set[str] = set()
    assigned: set[str] = set()
    pending: list[tuple[dict[str, Any], TranslationResult | None, str, str]] = []

    for row in country_rows:
        orig = row["orig_variable"]
        seen.add(orig)
        if orig in existing and not force_realign:
            saved = existing[orig]
            if saved.canonical_variable:
                assigned.add(saved.canonical_variable)
            alignment.append(saved)
            continue

        tr = (translations or {}).get(orig)
        canonical_variable, notes = _resolve_canonical_for_row(
            row,
            tr,
            core_whitelist,
            core_lookup,
            core_orig_map,
            assigned,
        )
        if canonical_variable:
            assigned.add(canonical_variable)
        pending.append((row, tr, canonical_variable, notes))

    canonical_counts: dict[str, int] = {}
    for _, _, cv, _ in pending:
        if cv:
            canonical_counts[cv] = canonical_counts.get(cv, 0) + 1
    for row in alignment:
        if row.canonical_variable:
            canonical_counts[row.canonical_variable] = (
                canonical_counts.get(row.canonical_variable, 0) + 1
            )

    for row, tr, canonical_variable, notes in pending:
        alignment.append(
            _row_from_parts(
                row,
                tr,
                canonical_variable,
                core_lookup,
                canonical_counts,
                notes=notes,
            )
        )

    for orig, saved in existing.items():
        if orig not in seen:
            alignment.append(saved)

    infer_scale_first_items(alignment, core_whitelist)
    for row in alignment:
        if row.canonical_variable and row.canonical_variable in core_lookup:
            cq, cit = core_lookup[row.canonical_variable]
            row.canonical_stem = cq.stem
            row.canonical_item = cit.item_text or ""
            row.canonical_options = _join_options(_option_labels(cq.options))
            if row.status == "extra":
                row.status = "probable"

    write_review_workbook(review_path, alignment, core_codebook, country_codebook)

    return alignment


def _ensure_orig_aliases(item: Any) -> dict[str, str]:
    aliases = getattr(item, "orig_aliases", None)
    if aliases is None:
        item.orig_aliases = {}
        return item.orig_aliases
    return aliases


def _alias_is_safe(alias: str, core_orig_ids_set: set[str], existing_index: dict[str, str]) -> bool:
    if alias in core_orig_ids_set and existing_index.get(alias.lower()) != alias:
        return False
    return True


def merge_country_alignment(
    core_codebook: list[Any],
    country_codebook: list[Any],
    alignment: list[AlignmentRow],
    country_source: str,
    unique_group_id_fn,
    assign_short_names_fn,
    used_gids: set[str],
    used_stems: set[str],
    used_names: set[str],
) -> list[Any]:
    """Apply alignment: aliases on core for matches, append extras for unmatched."""
    combined = list(core_codebook)
    core_lookup = _core_lookup(core_codebook)
    core_orig_ids_set = core_orig_ids(core_codebook)
    align_by_orig = {r.country_orig_id: r for r in alignment}

    country_q_by_item: dict[str, Any] = {}
    for q in country_codebook:
        for it in q.items:
            country_q_by_item[it.variable] = q

    matched_country_ids: set[str] = set()
    for row in alignment:
        canon = (row.canonical_variable or "").strip()
        if not canon or canon not in core_lookup:
            continue
        matched_country_ids.add(row.country_orig_id)
        cq, cit = core_lookup[canon]
        aliases = _ensure_orig_aliases(cit)
        aliases[country_source] = row.country_orig_id

        country_q = country_q_by_item.get(row.country_orig_id)
        if country_q and cq.options:
            country_opts = _option_labels(country_q.options)
            for idx, opt in enumerate(cq.options):
                if idx < len(country_opts):
                    label = country_opts[idx]
                    if label and label not in opt.aliases and label != opt.label:
                        opt.aliases.append(label)

    extras: list[Any] = []
    for q in country_codebook:
        new_items = [it for it in q.items if it.variable not in matched_country_ids]
        if not new_items:
            continue
        q_copy = deepcopy(q)
        q_copy.items = new_items
        q_copy.group_id = unique_group_id_fn(used_gids, q.stem, f"{country_source}_q")
        q_copy.source = country_source
        q_copy.is_core = False
        for it in q_copy.items:
            it.source = country_source
            it.is_core = False
        extras.append(q_copy)

    assign_short_names_fn(extras, used_stems, used_names)
    combined.extend(extras)
    return combined


def merge_by_shared_ids(
    core_codebook: list[Any],
    country_codebook: list[Any],
    country_source: str,
    core_vids: set[str],
    unique_group_id_fn,
    assign_short_names_fn,
    used_gids: set[str],
    used_stems: set[str],
    used_names: set[str],
) -> list[Any]:
    """Germany-style merge: keep only country variables whose IDs are not already in core."""
    combined = list(core_codebook)
    country_non_core: list[Any] = []
    for q in country_codebook:
        new_items = [it for it in q.items if it.variable not in core_vids]
        if not new_items:
            continue
        q_copy = deepcopy(q)
        q_copy.items = new_items
        q_copy.group_id = unique_group_id_fn(used_gids, q.stem, f"{country_source}_q")
        country_non_core.append(q_copy)

    assign_short_names_fn(country_non_core, used_stems, used_names)
    combined.extend(country_non_core)
    return combined
