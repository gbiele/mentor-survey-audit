#!/usr/bin/env python3
"""Tests for independent-country survey alignment (Spanish-case path)."""

import sys
import tempfile
import unittest
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from openpyxl import load_workbook

from country_alignment import (  # noqa: E402
    AlignmentRow,
    TranslationResult,
    align_country_to_core,
    fallback_match_canonical,
    id_overlap_ratio,
    infer_scale_first_items,
    load_review_workbook,
    merge_by_shared_ids,
    merge_country_alignment,
    sort_alignment_for_review,
    write_review_workbook,
    build_core_whitelist,
    _match_score,
)


@dataclass
class Option:
    label: str
    order: int = 0
    value: int | None = None
    aliases: list[str] = field(default_factory=list)


@dataclass
class Item:
    variable: str
    item_text: str = ""
    source: str = "canonical_en"
    short_name: str = ""
    is_core: bool = True
    orig_aliases: dict[str, str] = field(default_factory=dict)


@dataclass
class Question:
    group_id: str
    section: str
    stem: str
    question_type: str
    scale: str
    scale_confidence: str
    multiple: bool
    options: list[Option]
    items: list[Item]
    options_complete: str
    source: str
    is_core: bool = True
    notes: str = ""


def _used_sets(codebook):
    used_gids = {q.group_id for q in codebook}
    used_stems = {q.group_id for q in codebook}
    used_names = {it.short_name for q in codebook for it in q.items if it.short_name}
    return used_gids, used_stems, used_names


def _unique_group_id(used, stem, fallback):
    base = fallback
    gid = base
    n = 2
    while gid in used:
        gid = f"{base}_{n}"
        n += 1
    used.add(gid)
    return gid


def _assign_short_names(qs, used_stems=None, used_names=None):
    used_stems = used_stems or set()
    used_names = used_names or set()
    for q in qs:
        stem = q.group_id[:8] or "extra"
        for i, it in enumerate(q.items, 1):
            name = f"{stem}{i}"
            while name in used_names:
                name = f"{stem}_{i}"
            it.short_name = name
            used_names.add(name)
        used_stems.add(stem)


def _core_fixture():
    gender = Question(
        group_id="gender",
        section="Background",
        stem="What gender do you identify as?",
        question_type="single",
        scale="nominal",
        scale_confidence="high",
        multiple=False,
        options=[
            Option("Boy", 0, 1),
            Option("Girl", 1, 2),
            Option("Other gender identity", 2, 3),
            Option("Prefer not to answer", 3, 999),
        ],
        items=[Item("gender1", source="canonical_en", short_name="gender1")],
        options_complete="true",
        source="canonical_en",
        is_core=True,
    )
    grade = Question(
        group_id="grade",
        section="Background",
        stem="What grade are you in?",
        question_type="single",
        scale="ordinal",
        scale_confidence="high",
        multiple=False,
        options=[
            Option("3rd year of ESO", 0, 1),
            Option("4th year of ESO", 1, 2),
        ],
        items=[Item("grade1", source="canonical_en", short_name="grade1")],
        options_complete="true",
        source="canonical_en",
        is_core=True,
    )
    byear = Question(
        group_id="byear",
        section="Background",
        stem="What year were you born?",
        question_type="interval",
        scale="interval",
        scale_confidence="high",
        multiple=False,
        options=[Option(str(y), i, y) for i, y in enumerate(range(2010, 2022))],
        items=[Item("born1", source="canonical_en", short_name="byear1")],
        options_complete="true",
        source="canonical_en",
        is_core=True,
    )
    month = Question(
        group_id="month",
        section="Background",
        stem="What month were you born?",
        question_type="single",
        scale="ordinal",
        scale_confidence="high",
        multiple=False,
        options=[Option("January", 0, 1), Option("February", 1, 2)],
        items=[Item("ID289", source="canonical_en", short_name="bmonth1")],
        options_complete="true",
        source="canonical_en",
        is_core=True,
    )
    ks = Question(
        group_id="ks",
        section="Quality of life",
        stem="When you think about the past week...",
        question_type="matrix",
        scale="ordinal",
        scale_confidence="high",
        multiple=False,
        options=[
            Option("Never", 0, 0),
            Option("Rarely", 1, 1),
            Option("Sometimes", 2, 2),
        ],
        items=[Item("ks10", item_text="Have you been able to pay attention?", source="canonical_en", short_name="ks10")],
        options_complete="true",
        source="canonical_en",
        is_core=True,
    )
    sa1 = Question(
        group_id="bcfpi_sa",
        section="Mental health",
        stem="Please select the response that best reflects how you have felt or behaved over the past six months.",
        question_type="matrix",
        scale="ordinal",
        scale_confidence="high",
        multiple=False,
        options=[
            Option("NEVER true", 0, 0),
            Option("SOMETIMES true", 1, 1),
            Option("OFTEN true", 2, 2),
        ],
        items=[
            Item(
                "sa1",
                item_text="worry because your heart is racing or your palms are sweaty in social gatherings",
                source="canonical_en",
                short_name="bcfpi_sa1",
            )
        ],
        options_complete="true",
        source="canonical_en",
        is_core=True,
    )
    return [gender, grade, byear, month, ks, sa1]


def _spain_fixture():
    gender = Question(
        group_id="genero",
        section="Background",
        stem="¿Con qué género te identificas?",
        question_type="single",
        scale="nominal",
        scale_confidence="high",
        multiple=False,
        options=[
            Option("Chico", 0, 1),
            Option("Chica", 1, 2),
            Option("Otra identidad de género", 2, 3),
            Option("Prefiero no responder", 3, 999),
        ],
        items=[Item("genero1", source="spain", is_core=False)],
        options_complete="true",
        source="spain",
        is_core=False,
    )
    byear = Question(
        group_id="year",
        section="Background",
        stem="¿En qué año naciste?",
        question_type="interval",
        scale="interval",
        scale_confidence="high",
        multiple=False,
        options=[Option(str(y), i, y) for i, y in enumerate(range(2010, 2022))],
        items=[Item("ID20", source="spain", is_core=False)],
        options_complete="true",
        source="spain",
        is_core=False,
    )
    false_friend = Question(
        group_id="sa",
        section="Mental health",
        stem="Regular la ansiedad social. Durante los últimos seis meses...",
        question_type="matrix",
        scale="ordinal",
        scale_confidence="high",
        multiple=False,
        options=[Option("Nunca", 0, 0), Option("A veces", 1, 1), Option("A menudo", 2, 2)],
        items=[Item("ID289", item_text="Te preocupas en situaciones sociales.", source="spain", is_core=False)],
        options_complete="true",
        source="spain",
        is_core=False,
    )
    cssrs = Question(
        group_id="cssrs",
        section="Suicide",
        stem="¿Cuántas veces has intentado suicidarte?",
        question_type="single",
        scale="ordinal",
        scale_confidence="high",
        multiple=False,
        options=[Option("Nunca", 0, 0), Option("Una vez", 1, 1)],
        items=[Item("ID829", source="spain", is_core=False)],
        options_complete="true",
        source="spain",
        is_core=False,
    )
    return [gender, byear, false_friend, cssrs]


def _germany_like_fixture(core):
    return [
        Question(
            group_id=q.group_id,
            section=q.section,
            stem=q.stem,
            question_type=q.question_type,
            scale=q.scale,
            scale_confidence=q.scale_confidence,
            multiple=q.multiple,
            options=list(q.options),
            items=[Item(it.variable, it.item_text, source="germany", is_core=False) for it in q.items],
            options_complete=q.options_complete,
            source="germany",
            is_core=False,
        )
        for q in core
    ]


class TestCountryAlignment(unittest.TestCase):
    def test_id_overlap_selects_merge_path(self):
        core = _core_fixture()
        spain = _spain_fixture()
        germany = _germany_like_fixture(core)
        self.assertLess(id_overlap_ratio(spain, core), 0.4)
        self.assertGreaterEqual(id_overlap_ratio(germany, core), 0.4)

    def test_merge_spanish_aliases_not_extras(self):
        core = _core_fixture()
        spain = _spain_fixture()
        translations = {
            "genero1": TranslationResult("genero1", canonical_id="gender1"),
            "ID20": TranslationResult("ID20", canonical_id="byear1"),
            "ID289": TranslationResult("ID289", canonical_id=None),
            "ID829": TranslationResult("ID829", canonical_id=None),
        }
        with tempfile.TemporaryDirectory() as tmp:
            review_path = Path(tmp) / "spain_review.xlsx"
            alignment = align_country_to_core(
                spain, core, review_path, "spain", translations=translations
            )
            used_gids, used_stems, used_names = _used_sets(core)
            merged = merge_country_alignment(
                core,
                spain,
                alignment,
                "spain",
                _unique_group_id,
                _assign_short_names,
                used_gids,
                used_stems,
                used_names,
            )

        gender_item = next(it for q in merged if q.is_core for it in q.items if it.short_name == "gender1")
        byear_item = next(it for q in merged if q.is_core for it in q.items if it.short_name == "byear1")
        self.assertEqual(gender_item.orig_aliases.get("spain"), "genero1")
        self.assertEqual(byear_item.orig_aliases.get("spain"), "ID20")

        extras = [it for q in merged if not q.is_core for it in q.items]
        extra_ids = {it.variable for it in extras}
        self.assertIn("ID829", extra_ids)
        self.assertIn("ID289", extra_ids)
        self.assertNotIn("genero1", extra_ids)
        self.assertNotIn("ID20", extra_ids)

    def test_false_friend_ids_not_merged_to_core_by_id(self):
        core = _core_fixture()
        spain = _spain_fixture()
        translations = {
            "genero1": TranslationResult("genero1", canonical_id="gender1"),
            "ID20": TranslationResult("ID20", canonical_id="byear1"),
            "ID289": TranslationResult(
                "ID289",
                translated_stem="Regulating social anxiety. During the past six months...",
                translated_item="You worry or feel very embarrassed in social situations.",
                translated_options=["Never true", "Sometimes true", "Often true"],
                canonical_id="bmonth1",
            ),
            "ID829": TranslationResult("ID829", canonical_id=None),
        }
        with tempfile.TemporaryDirectory() as tmp:
            review_path = Path(tmp) / "spain_review.xlsx"
            alignment = align_country_to_core(
                spain, core, review_path, "spain", translations=translations
            )
            id289_row = next(r for r in alignment if r.country_orig_id == "ID289")
            self.assertEqual(id289_row.canonical_variable, "bcfpi_sa1")
            self.assertNotEqual(id289_row.canonical_variable, "bmonth1")

            used_gids, used_stems, used_names = _used_sets(core)
            merged = merge_country_alignment(
                core,
                spain,
                alignment,
                "spain",
                _unique_group_id,
                _assign_short_names,
                used_gids,
                used_stems,
                used_names,
            )
            bmonth = next(
                (it for q in merged if q.is_core for it in q.items if it.short_name == "bmonth1"),
                None,
            )
            self.assertIsNotNone(bmonth)
            self.assertNotIn("spain", bmonth.orig_aliases)
            sa1 = next(
                (it for q in merged if q.is_core for it in q.items if it.short_name == "bcfpi_sa1"),
                None,
            )
            self.assertEqual(sa1.orig_aliases.get("spain"), "ID289")

    def test_fallback_match_when_llm_leaves_canonical_empty(self):
        core = _core_fixture()
        whitelist = build_core_whitelist(core)
        row = {
            "orig_variable": "grado1",
            "country_stem": "¿En qué curso estás actualmente?",
            "country_item": "",
            "country_options": ["3rd year", "4th year"],
        }
        tr = TranslationResult(
            "grado1",
            translated_stem="What grade are you in?",
            translated_options=["3rd year of ESO", "4th year of ESO"],
            canonical_id=None,
        )
        row["country_options"] = ["3rd year of ESO", "4th year of ESO"]
        var, note = fallback_match_canonical(tr, row, whitelist, set(), {}, "grado1")
        self.assertEqual(var, "grade1")
        self.assertIn("Fallback", note)

    def test_grade_match_tolerates_extra_country_options(self):
        core = _core_fixture()
        whitelist = build_core_whitelist(core)
        grade_entry = next(c for c in whitelist if c["variable"] == "grade1")
        tr = TranslationResult(
            "grado1",
            translated_stem="What grade are you currently in?",
            canonical_id=None,
        )
        spain_options = [f"Option {i}" for i in range(10)]
        score = _match_score(tr, "¿En qué curso estás?", "", spain_options, grade_entry)
        self.assertGreaterEqual(score, 0.55)

        row = {
            "orig_variable": "grado1",
            "country_stem": "¿En qué curso estás?",
            "country_item": "",
            "country_options": spain_options,
        }
        var, _ = fallback_match_canonical(tr, row, whitelist, set(), {}, "grado1")
        self.assertEqual(var, "grade1")

    def test_matrix_match_uses_country_item_when_translation_item_empty(self):
        core = _core_fixture()
        whitelist = build_core_whitelist(core)
        ks_entry = next(c for c in whitelist if c["variable"] == "ks10")
        tr = TranslationResult(
            "ks2",
            translated_stem="When you think about the past week...",
            translated_item="",
            canonical_id=None,
        )
        country_item = "Have you felt sad?"
        core_item = ks_entry["item_text"]
        ks2_core = dict(ks_entry)
        ks2_core["variable"] = "ks2"
        ks2_core["item_text"] = "Have you felt sad?"
        spain_options = [f"Scale {i}" for i in range(8)]
        score = _match_score(
            tr,
            "Durante la última semana...",
            country_item,
            spain_options,
            ks2_core,
        )
        self.assertGreaterEqual(score, 0.42)

    def test_scale_inference_does_not_cross_prefix(self):
        core_whitelist = [
            {
                "variable": "dlup2",
                "section": "ACE",
                "question_stem": "During your life...",
                "item_text": "Often felt that no one loved you",
                "options": ["Never", "Once", "More than once"],
            },
            {
                "variable": "rpginsfy3",
                "section": "Family",
                "question_stem": "During your life...",
                "item_text": "Lived with someone who was depressed",
                "options": ["Yes", "No"],
            },
            {
                "variable": "rpginsfy1",
                "section": "Family",
                "question_stem": "Who is your parent or guardian?",
                "item_text": "",
                "options": ["Mother", "Father", "Other"],
            },
        ]
        alignment = [
            AlignmentRow(
                country_orig_id="ID392",
                section="Family",
                country_stem="Durante tu vida...",
                country_item="¿Viviste con alguien deprimido?",
                country_options="Si | No",
                translated_text="During your life... : Lived with someone depressed | Yes | No",
                canonical_variable="",
                canonical_stem="",
                canonical_item="",
                canonical_options="",
                status="extra",
            ),
            AlignmentRow(
                country_orig_id="ID400",
                section="ACE",
                country_stem="During your life...",
                country_item="Often felt unloved",
                country_options="Never | Once | More than once",
                translated_text="During your life... : Often felt unloved | Never | Once | More than once",
                canonical_variable="dlup2",
                canonical_stem="During your life...",
                canonical_item="Often felt that no one loved you",
                canonical_options="Never | Once | More than once",
                status="matched",
            ),
            AlignmentRow(
                country_orig_id="ID401",
                section="ACE",
                country_stem="During your life...",
                country_item="Lived with depressed person",
                country_options="Yes | No",
                translated_text="During your life... : Lived with depressed person | Yes | No",
                canonical_variable="rpginsfy3",
                canonical_stem="During your life...",
                canonical_item="Lived with someone who was depressed",
                canonical_options="Yes | No",
                status="matched",
            ),
        ]
        infer_scale_first_items(alignment, core_whitelist)
        id392 = next(r for r in alignment if r.country_orig_id == "ID392")
        self.assertNotEqual(id392.canonical_variable, "dlup2")
        self.assertIn(id392.canonical_variable, ("", "rpginsfy1"))

    def test_education_match_tolerates_extra_options(self):
        whitelist = [
            {
                "variable": "edumom1",
                "section": "Background",
                "question_stem": "What best describes your mother's education?",
                "item_text": "",
                "options": ["Primary", "Secondary", "College", "Other", "Unknown", "NA"],
            },
            {
                "variable": "edudad1",
                "section": "Background",
                "question_stem": "What best describes your father's education?",
                "item_text": "",
                "options": ["Primary", "Secondary", "College", "Other", "Unknown", "NA"],
            },
        ]
        row = {
            "orig_variable": "ID62",
            "country_stem": "¿Cuál es el nivel educativo más alto alcanzado por tu madre?",
            "country_item": "",
            "country_options": [f"Level {i}" for i in range(11)],
        }
        tr = TranslationResult(
            "ID62",
            translated_stem="What is the highest level of education attained by your mother?",
            canonical_id=None,
        )
        var, note = fallback_match_canonical(tr, row, whitelist, set(), {}, "ID62")
        self.assertEqual(var, "edumom1")
        self.assertIn("Fallback", note)

    def test_exact_variable_name_match_for_ks2(self):
        whitelist = [
            {
                "variable": "ks2",
                "section": "QoL",
                "question_stem": "When you think about the past week...",
                "item_text": "Have you felt sad?",
                "options": ["Never", "Rarely", "Sometimes", "Often", "Always"],
            }
        ]
        row = {
            "orig_variable": "ks2",
            "country_stem": "Pensando en la última semana...",
            "country_item": "",
            "country_options": [f"Opt {i}" for i in range(8)],
        }
        tr = TranslationResult(
            "ks2",
            translated_stem="Thinking about the past week...",
            canonical_id=None,
        )
        var, note = fallback_match_canonical(tr, row, whitelist, set(), {}, "ks2")
        self.assertEqual(var, "ks2")
        self.assertIn("Exact variable name", note)

    def test_bage1_preferred_over_bmborn1_for_combo_question(self):
        core_whitelist = [
            {
                "variable": "bage1",
                "section": "Background",
                "question_stem": "How old were you when you came to this country?",
                "item_text": "",
                "options": [str(i) for i in range(1, 19)],
            },
            {
                "variable": "bmborn1",
                "section": "Background",
                "question_stem": "In which country were you born?",
                "item_text": "",
                "options": ["Spain", "Other country"],
            },
        ]
        row = {
            "orig_variable": "ID67",
            "country_stem": "¿Qué edad tenías cuando llegaste a este país y de qué país venías?",
            "country_item": "",
            "country_options": ["0-5", "6-10", "Spain", "Other"],
        }
        tr = TranslationResult(
            "ID67",
            translated_stem="How old were you when you arrived in this country and from which country?",
            canonical_id=None,
        )
        var, note = fallback_match_canonical(tr, row, core_whitelist, set(), {}, "ID67")
        self.assertEqual(var, "bage1")
        self.assertIn("Fallback", note)

    def test_review_workbook_follows_questionnaire_order(self):
        core = _core_fixture()
        spain = [
            Question(
                group_id="genero",
                section="Background",
                stem="Gender",
                question_type="single",
                scale="nominal",
                scale_confidence="high",
                multiple=False,
                options=[Option("Boy", 0, 1)],
                items=[Item("genero1", source="spain", is_core=False)],
                options_complete="true",
                source="spain",
                is_core=False,
            ),
            Question(
                group_id="extra_sleep",
                section="Background",
                stem="Sleep extra",
                question_type="single",
                scale="ordinal",
                scale_confidence="high",
                multiple=False,
                options=[Option("Never", 0, 0)],
                items=[Item("ID999", source="spain", is_core=False)],
                options_complete="true",
                source="spain",
                is_core=False,
            ),
            Question(
                group_id="year",
                section="Background",
                stem="Birth year",
                question_type="interval",
                scale="interval",
                scale_confidence="high",
                multiple=False,
                options=[Option("2010", 0, 2010)],
                items=[Item("ID20", source="spain", is_core=False)],
                options_complete="true",
                source="spain",
                is_core=False,
            ),
        ]
        translations = {
            "genero1": TranslationResult("genero1", canonical_id="gender1"),
            "ID20": TranslationResult("ID20", canonical_id="byear1"),
        }
        with tempfile.TemporaryDirectory() as tmp:
            review_path = Path(tmp) / "spain_review.xlsx"
            align_country_to_core(spain, core, review_path, "spain", translations=translations)

            wb = load_workbook(review_path)
            ws = wb.active
            headers = [c.value for c in ws[1]]
            orig_idx = headers.index("country_orig_id")
            status_idx = headers.index("status")
            canon_idx = headers.index("canonical_variable")

            sheet_order = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[status_idx] == "missing_core":
                    sheet_order.append(f"missing:{row[canon_idx]}")
                else:
                    sheet_order.append(row[orig_idx])

            self.assertEqual(sheet_order[0], "genero1")
            self.assertEqual(sheet_order[1], "ID999")
            self.assertIn("missing:grade1", sheet_order)
            self.assertLess(sheet_order.index("ID999"), sheet_order.index("ID20"))
            self.assertLess(sheet_order.index("genero1"), sheet_order.index("ID999"))

    def test_sort_alignment_interleaves_extras_between_core_slots(self):
        core = _core_fixture()[:2]  # gender1, grade1
        spain = [
            Question(
                group_id="genero",
                section="Background",
                stem="Gender",
                question_type="single",
                scale="nominal",
                scale_confidence="high",
                multiple=False,
                options=[Option("Boy", 0, 1)],
                items=[Item("genero1", source="spain", is_core=False)],
                options_complete="true",
                source="spain",
                is_core=False,
            ),
            Question(
                group_id="extra",
                section="Background",
                stem="Extra",
                question_type="single",
                scale="nominal",
                scale_confidence="high",
                multiple=False,
                options=[Option("Yes", 0, 1)],
                items=[Item("ID999", source="spain", is_core=False)],
                options_complete="true",
                source="spain",
                is_core=False,
            ),
            Question(
                group_id="grade",
                section="Background",
                stem="Grade",
                question_type="single",
                scale="ordinal",
                scale_confidence="high",
                multiple=False,
                options=[Option("3rd", 0, 1)],
                items=[Item("grado1", source="spain", is_core=False)],
                options_complete="true",
                source="spain",
                is_core=False,
            ),
        ]
        rows = [
            AlignmentRow(
                country_orig_id="grado1",
                section="Background",
                country_stem="Grade",
                country_item="",
                country_options="",
                translated_text="",
                canonical_variable="grade1",
                canonical_stem="",
                canonical_item="",
                canonical_options="",
                status="matched",
            ),
            AlignmentRow(
                country_orig_id="genero1",
                section="Background",
                country_stem="Gender",
                country_item="",
                country_options="",
                translated_text="",
                canonical_variable="gender1",
                canonical_stem="",
                canonical_item="",
                canonical_options="",
                status="matched",
            ),
            AlignmentRow(
                country_orig_id="ID999",
                section="Background",
                country_stem="Extra",
                country_item="",
                country_options="",
                translated_text="",
                canonical_variable="",
                canonical_stem="",
                canonical_item="",
                canonical_options="",
                status="extra",
            ),
        ]
        ordered = sort_alignment_for_review(rows, core, spain)
        self.assertEqual(
            [r.country_orig_id or f"missing:{r.canonical_variable}" for r in ordered],
            ["genero1", "ID999", "grado1"],
        )

    def test_review_workbook_round_trip(self):
        core = _core_fixture()
        spain = _spain_fixture()
        translations = {
            "genero1": TranslationResult("genero1", canonical_id="gender1"),
            "ID20": TranslationResult("ID20", canonical_id=None),
            "ID289": TranslationResult("ID289", canonical_id=None),
            "ID829": TranslationResult("ID829", canonical_id=None),
        }
        with tempfile.TemporaryDirectory() as tmp:
            review_path = Path(tmp) / "spain_review.xlsx"
            align_country_to_core(spain, core, review_path, "spain", translations=translations)

            loaded = load_review_workbook(review_path)
            loaded["ID20"].canonical_variable = "byear1"
            loaded["ID20"].status = "matched"
            write_review_workbook(review_path, list(loaded.values()), core)

            alignment = align_country_to_core(spain, core, review_path, "spain", translations={})
            id20 = next(r for r in alignment if r.country_orig_id == "ID20")
            self.assertEqual(id20.canonical_variable, "byear1")

            used_gids, used_stems, used_names = _used_sets(core)
            merged = merge_country_alignment(
                core,
                spain,
                alignment,
                "spain",
                _unique_group_id,
                _assign_short_names,
                used_gids,
                used_stems,
                used_names,
            )
            byear_item = next(it for q in merged if q.is_core for it in q.items if it.short_name == "byear1")
            self.assertEqual(byear_item.orig_aliases.get("spain"), "ID20")

    def test_review_workbook_maps_one_extra_to_two_core_ids(self):
        core = _core_fixture()
        rpg = Question(
            group_id="rpginsfy",
            section="Adverse childhood experiences",
            stem="Relationship with parents/guardians",
            question_type="matrix",
            scale="ordinal",
            scale_confidence="high",
            multiple=False,
            options=[Option("Always", 0, 4), Option("Never", 1, 0)],
            items=[
                Item("ace01", item_text="understood problems", source="canonical_en", short_name="rpginsfy1"),
                Item("ace02", item_text="knew free time", source="canonical_en", short_name="rpginsfy2"),
            ],
            options_complete="true",
            source="canonical_en",
            is_core=True,
        )
        core = list(core) + [rpg]
        spain = [
            Question(
                group_id="id391",
                section="Adverse childhood experiences",
                stem="Relación con los padres/tutores",
                question_type="matrix",
                scale="ordinal",
                scale_confidence="high",
                multiple=False,
                options=[Option("Siempre", 0, 4), Option("Nunca", 1, 0)],
                items=[Item("ID391", source="spain", is_core=False)],
                options_complete="true",
                source="spain",
                is_core=False,
            )
        ]
        with tempfile.TemporaryDirectory() as tmp:
            review_path = Path(tmp) / "spain_review.xlsx"
            write_review_workbook(
                review_path,
                [
                    AlignmentRow(
                        country_orig_id="ID391",
                        section="Adverse childhood experiences",
                        country_stem="Relación con los padres/tutores",
                        country_item="",
                        country_options="Siempre | Nunca",
                        translated_text="",
                        canonical_variable="rpginsfy1, rpginsfy2",
                        canonical_stem="",
                        canonical_item="",
                        canonical_options="",
                        status="probable",
                        notes="Manual review",
                    )
                ],
            )
            alignment = align_country_to_core(spain, core, review_path, "spain", translations={})
            mapped = {
                r.canonical_variable: r.country_orig_id
                for r in alignment
                if r.canonical_variable in {"rpginsfy1", "rpginsfy2"}
            }
            self.assertEqual(mapped["rpginsfy1"], "ID391")
            self.assertEqual(mapped["rpginsfy2"], "ID391")
            wb = load_workbook(review_path)
            ws = wb.active
            headers = [c.value for c in ws[1]]
            status_idx = headers.index("status")
            canon_idx = headers.index("canonical_variable")
            missing = {
                row[canon_idx]
                for row in ws.iter_rows(min_row=2, values_only=True)
                if row[status_idx] == "missing_core"
            }
            self.assertNotIn("rpginsfy1", missing)
            self.assertNotIn("rpginsfy2", missing)

            used_gids, used_stems, used_names = _used_sets(core)
            merged = merge_country_alignment(
                core,
                spain,
                alignment,
                "spain",
                _unique_group_id,
                _assign_short_names,
                used_gids,
                used_stems,
                used_names,
            )
            by_name = {it.short_name: it for q in merged for it in q.items}
            self.assertEqual(by_name["rpginsfy1"].orig_aliases.get("spain"), "ID391")
            self.assertEqual(by_name["rpginsfy2"].orig_aliases.get("spain"), "ID391")

    def test_review_workbook_shows_missing_core_in_red(self):
        core = _core_fixture()
        spain = _spain_fixture()
        translations = {
            "genero1": TranslationResult("genero1", canonical_id="gender1"),
            "ID20": TranslationResult("ID20", canonical_id=None),
            "ID289": TranslationResult("ID289", canonical_id=None),
            "ID829": TranslationResult("ID829", canonical_id=None),
        }
        with tempfile.TemporaryDirectory() as tmp:
            review_path = Path(tmp) / "spain_review.xlsx"
            align_country_to_core(spain, core, review_path, "spain", translations=translations)

            wb = load_workbook(review_path)
            ws = wb.active
            headers = [c.value for c in ws[1]]
            status_idx = headers.index("status")
            canon_idx = headers.index("canonical_variable")

            missing_rows = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[status_idx] == "missing_core":
                    missing_rows.append(row[canon_idx])

            self.assertIn("byear1", missing_rows)
            self.assertIn("bmonth1", missing_rows)
            self.assertIn("ks10", missing_rows)
            self.assertNotIn("gender1", missing_rows)

            status_col = status_idx + 1
            for row_num in range(2, ws.max_row + 1):
                if ws.cell(row_num, status_col).value == "missing_core":
                    fill = ws.cell(row_num, 1).fill.start_color.rgb
                    self.assertIn("FF6666", str(fill))

    def test_merge_by_shared_ids_for_high_overlap(self):
        core = _core_fixture()
        germany = _germany_like_fixture(core)
        core_vids = {it.variable for q in core for it in q.items}
        used_gids, used_stems, used_names = _used_sets(core)
        merged = merge_by_shared_ids(
            core,
            germany,
            "germany",
            core_vids,
            _unique_group_id,
            _assign_short_names,
            used_gids,
            used_stems,
            used_names,
        )
        self.assertEqual(len(merged), len(core))


if __name__ == "__main__":
    unittest.main()
